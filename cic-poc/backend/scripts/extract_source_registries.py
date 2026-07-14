#!/usr/bin/env python3
"""One-time script to extract Source Registry workbooks into JSON for the backend."""

import json
import re
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).parent.parent.parent.parent


def snake_case(header: str) -> str:
    header = header.strip().replace("/", " ").replace("-", " ")
    header = re.sub(r"[^\w\s]", "", header)
    return re.sub(r"\s+", "_", header.strip()).lower()


def extract_sheet(path: Path, sheet_name: str, id_column: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    header_row = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    keys = [snake_case(h) if h else f"col_{i}" for i, h in enumerate(header_row)]
    id_index = keys.index(snake_case(id_column))

    rows = []
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if values[id_index] in (None, ""):
            continue
        row = {}
        for key, value in zip(keys, values):
            if key == keys[id_index]:
                row["id"] = str(value)
            else:
                row[key] = value if value is not None else ""
        rows.append(row)
    return rows


def main():
    jobs = [
        (
            REPO_ROOT / "Syriac-Build/World-Builds/Syriac-Christianity-Edessa-Nisibis/Source_Registry.xlsx",
            "Registry",
            "#",
            REPO_ROOT / "cic-poc/backend/data/syriac_world/source_registry.json",
        ),
        (
            REPO_ROOT / "World-Builds/01-Post-Apostolic-House-Church/CiC_W1_Source_Registry_FINAL_v2.xlsx",
            "Source Registry",
            "ID",
            REPO_ROOT / "cic-poc/backend/data/pahc_world/source_registry.json",
        ),
    ]

    for src_path, sheet_name, id_column, out_path in jobs:
        rows = extract_sheet(src_path, sheet_name, id_column)
        out_path.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"{src_path.name} [{sheet_name}] -> {out_path} ({len(rows)} rows)")
        print(f"  columns: {list(rows[0].keys())}")
        print(f"  row[0]: {json.dumps(rows[0], ensure_ascii=False)[:300]}")
        print(f"  row[1]: {json.dumps(rows[1], ensure_ascii=False)[:300]}")


if __name__ == "__main__":
    main()
